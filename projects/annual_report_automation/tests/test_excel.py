"""Excel 출력 검증. 근거: plans/PLAN-4-세무조정.md 4-10 (근거 각주 붙은 세무사 전달용 초안)."""

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from taxengine.pipeline import 실행
from taxengine.export.excel import 작성

FIXTURE = Path(__file__).resolve().parent.parent / "data" / "example" / "fy2025"
CAR_FIXTURE = Path(__file__).resolve().parent / "fixtures" / "car-demo"


class 작성_테스트(unittest.TestCase):
    def test_기본_시트_구성(self):
        out = 실행(FIXTURE)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "test.xlsx"
            작성(out, path)
            wb = load_workbook(path)
            self.assertIn("별지3", wb.sheetnames)
            self.assertIn("감가상각시부인", wb.sheetnames)
            self.assertIn("기업업무추진비", wb.sheetnames)
            self.assertIn("정답대조", wb.sheetnames)
            self.assertNotIn("업무용승용차", wb.sheetnames)  # fy2025엔 cars.csv 없음

    def test_별지3_시트에_차감납부세액과_근거가_있다(self):
        out = 실행(FIXTURE)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "test.xlsx"
            작성(out, path)
            ws = load_workbook(path)["별지3"]
            라벨들 = [row[1].value for row in ws.iter_rows(min_row=2)]
            self.assertIn("차감 납부할 세액", 라벨들)
            마지막행 = list(ws.iter_rows(min_row=2))[-1]
            self.assertEqual(마지막행[2].value, int(out["r"]["차감납부세액"]))
            self.assertTrue(마지막행[3].value)  # 근거 칸이 비어있지 않음

    def test_모든_금액줄에_근거가_비어있지_않다(self):
        # 이월결손금 공제처럼 라벨이 동적으로 바뀌는 줄도 근거를 빠뜨리면 안 된다
        out = 실행(FIXTURE)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "test.xlsx"
            작성(out, path)
            ws = load_workbook(path)["별지3"]
            for row in ws.iter_rows(min_row=2):
                라벨, 근거 = row[1].value, row[3].value
                if 라벨 == "세율":
                    continue  # 세율 줄은 금액이 아니라 구간 문구가 근거 칸에 이미 들어감
                self.assertTrue(근거, f"'{라벨}' 줄에 근거가 비어있음")

    def test_정답대조_시트에_일치_표시(self):
        out = 실행(FIXTURE)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "test.xlsx"
            작성(out, path)
            ws = load_workbook(path)["정답대조"]
            일치칸들 = [row[3].value for row in ws.iter_rows(min_row=2)]
            self.assertTrue(all("일치" in v for v in 일치칸들))
            self.assertTrue(all(v.startswith("✓") for v in 일치칸들))  # fy2025는 5/5 재현 성공

    def test_cars_csv_있으면_업무용승용차_시트가_생긴다(self):
        out = 실행(CAR_FIXTURE)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "test.xlsx"
            작성(out, path)
            wb = load_workbook(path)
            self.assertIn("업무용승용차", wb.sheetnames)
            ws = wb["업무용승용차"]
            self.assertEqual(ws.cell(row=2, column=1).value, "차량1")


if __name__ == "__main__":
    unittest.main()
