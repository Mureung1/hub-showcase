"""Excel 출력 — 세무사 전달용 초안. 별지3·감가상각시부인·기업업무추진비·업무용승용차·정답대조를
근거 각주와 함께 시트로 나눠 담는다.

근거: plans/PLAN-4-세무조정.md 4-10 ("수치마다 근거 각주 부착 → 세무사 전달용 초안의 요건")
"""

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from taxengine.money import 원

굵게 = Font(bold=True)
숫자서식 = "#,##0"

# 별지3 각 라벨의 근거 — r["단계"]에 이미 "주석"이 붙은 줄(세율 구간 등)은 그걸 우선 쓰고,
# 없는 줄만 여기서 보완한다.
별지3_각주 = {
    "결산서상 당기순이익": "손익계산서: 수익 합계 − 비용 합계 (법인세비용 포함)",
    "익금산입·손금불산입": "소득금액조정합계표 + 감가상각부인액 + 기업업무추진비 한도초과 + 업무용승용차 한도초과",
    "손금산입·익금불산입": "소득금액조정합계표 + 감가상각 시인부족 추인액",
    "기부금 한도초과": "company.csv 기부금한도초과 — 별지3에서 직접 가감(소득금액조정합계표엔 안 넣음)",
    "각 사업연도 소득금액": "당기순이익 ± 조정 (법인세법 §14)",
    "과세표준": "각사업연도소득 − 이월결손금공제 (법인세법 §13)",
    "산출세액": "과세표준 × 누진세율 (법인세법 §55)",
    "공제·감면세액": "company.csv 공제감면세액 (최저한세 배제 반영 후 금액)",
    "기납부세액": "company.csv 기납부세액 (원천납부세액명세서 + 중간예납)",
    "차감 납부할 세액": "산출세액 − 공제감면세액 + 가산세 − 기납부세액 (최저한세 배제 반영)",
}


def _헤더(ws, 헤더들):
    for i, h in enumerate(헤더들, 1):
        ws.cell(row=1, column=i, value=h).font = 굵게
    ws.freeze_panes = "A2"


def _열너비(ws, 너비들):
    for i, w in enumerate(너비들, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _숫자열서식(ws, 열번호들):
    for col in 열번호들:
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = 숫자서식


def _별지3(wb, out):
    ws = wb.active
    ws.title = "별지3"
    _헤더(ws, ["부호", "항목", "금액", "근거·산식"])
    for s in out["r"]["단계"]:
        금액 = int(s["금액"]) if s["금액"] is not None else None
        라벨 = s["라벨"]
        if 라벨.startswith("이월결손금 공제"):
            근거 = "이월결손금 잔액과 (각사업연도소득 × 한도율) 중 작은 값 (법인세법 §13)"
        elif 라벨.startswith("최저한세 배제"):
            근거 = "조특법 §132 — 공제감면 후 세액이 최저한세에 못 미치는 부족분만큼 배제"
        else:
            근거 = s.get("주석") or 별지3_각주.get(라벨, "")
        ws.append([s["부호"], s["라벨"], 금액, 근거])
        if s.get("강조") or s.get("최종"):
            for cell in ws[ws.max_row]:
                cell.font = 굵게
    _숫자열서식(ws, [3])
    _열너비(ws, [4, 26, 16, 64])


def _감가상각(wb, out):
    시부인들 = out.get("시부인들") or []
    if not 시부인들:
        return
    ws = wb.create_sheet("감가상각시부인")
    _헤더(ws, ["자산", "방법", "내용연수", "회사계상액", "상각범위액", "판정", "부인액", "추인액", "산식(근거)"])
    for s in 시부인들:
        ws.append([s["자산"], s["방법"], s["내용연수"], int(s["회사계상액"]), int(s["상각범위액"]),
                   s["판정"], int(s["부인액"]), int(s["추인액"]), s["산식"]])
    _숫자열서식(ws, [4, 5, 7, 8])
    _열너비(ws, [12, 8, 8, 14, 14, 26, 12, 12, 56])


def _추진비(wb, out):
    추진비 = out.get("추진비")
    if not 추진비:
        return
    ws = wb.create_sheet("기업업무추진비")
    _헤더(ws, ["항목", "금액", "근거"])
    rows = [
        ("회사계상액", 추진비["회사계상액"], "손익계산서 기업업무추진비 계정"),
        ("적격증빙없는금액", 추진비["적격증빙없는금액"],
         "company.csv 기업업무추진비_증빙불비금액 (법인세법 §25② 적격증빙 미수취분, 한도와 무관하게 전액 부인)"),
        ("적격지출", 추진비["적격지출"], "회사계상액 − 적격증빙없는금액"),
        ("한도", 추진비["한도"], "기본한도×월수/12 + 수입금액×구간별 적용률 (소규모임대법인이면 50% 축소, 법인세법 §25)"),
        ("한도초과", 추진비["한도초과"], "적격지출 − 한도 (0 미만이면 0)"),
        ("손금불산입합계", 추진비["손금불산입합계"], "적격증빙없는금액 + 한도초과 → 별지3 가산조정에 자동 합산"),
    ]
    for label, amt, note in rows:
        ws.append([label, int(amt), note])
    _숫자열서식(ws, [2])
    _열너비(ws, [20, 16, 70])


def _승용차(wb, out):
    차량판정들 = out.get("차량판정들") or []
    차량대장 = out.get("차량대장") or []
    if not 차량판정들:
        return
    ws = wb.create_sheet("업무용승용차")
    _헤더(ws, ["차량", "관련비용합계", "업무사용비율", "손금인정액", "감가상각한도초과", "손금불산입액", "사유"])
    for 차량, c in zip(차량대장, 차량판정들):
        ws.append([차량.get("명", ""), int(c["관련비용합계"]), f"{float(c['업무사용비율']) * 100:.0f}%",
                   int(c["손금인정액"]), int(c["감가상각한도초과"]), int(c["손금불산입액"]), c["사유"]])
    _숫자열서식(ws, [2, 4, 5, 6])
    _열너비(ws, [12, 14, 12, 14, 16, 14, 50])


def _정답대조(wb, out):
    정답 = out.get("정답")
    if not 정답:
        return
    r = out["r"]
    후보 = [
        ("각사업연도소득", r["각사업연도소득"], 정답.get("각사업연도소득")),
        ("과세표준", r["과세표준"], 정답.get("과세표준")),
        ("산출세액", r["산출세액"], 정답.get("산출세액")),
        ("차감납부세액", r["차감납부세액"], 정답.get("차감납부세액")),
        ("지방소득세", r["지방소득세"]["산출세액"], 정답.get("지방소득세")),
    ]
    행들 = [(n, e, a) for n, e, a in 후보 if a is not None]
    if not 행들:
        return

    ws = wb.create_sheet("정답대조")
    _헤더(ws, ["항목", "엔진", "실제(종이책)", "일치"])
    for name, 엔진, 실제 in 행들:
        일치 = 원(엔진) == 원(실제)
        ws.append([name, int(엔진), int(원(실제)), "✓ 일치" if 일치 else "✗ 불일치"])
    _숫자열서식(ws, [2, 3])
    _열너비(ws, [18, 16, 16, 10])


def 작성(out: dict, path) -> None:
    """pipeline.실행() 결과를 근거 각주가 붙은 xlsx로 저장한다."""
    wb = Workbook()
    _별지3(wb, out)
    _감가상각(wb, out)
    _추진비(wb, out)
    _승용차(wb, out)
    _정답대조(wb, out)
    wb.save(path)
